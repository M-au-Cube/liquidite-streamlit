"""
Construction des deux documents Word : format ArXiv et format BMI (gabarit de l'ancienne version).
Usage : python build_docx.py
"""
import copy, os, re, sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "tools"))
from omml import latex_to_omml  # noqa
import content as C  # noqa

FIG = os.path.join(HERE, "..", "figs")
OUT = os.path.join(HERE, "..", "livrables")
os.makedirs(OUT, exist_ok=True)
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


# ---------------------------------------------------------------------------
# utilitaires
# ---------------------------------------------------------------------------
def eq_numbers(blocks_list):
    num, n = {}, 0
    for blocks in blocks_list:
        for b in blocks:
            if b[0] == "eq":
                n += 1
                if b[2]:
                    num[b[2]] = n
    return num

EQN = eq_numbers([C.B, C.APP])

def resolve_refs(text):
    text = re.sub(r"\(équation \[\[(.+?)\]\]\)", lambda m: f"(équation {EQN[m.group(1)]})", text)
    return re.sub(r"\[\[(.+?)\]\]", lambda m: f"({EQN[m.group(1)]})", text)

def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kw:
            el = borders.find(qn("w:" + edge))
            if el is None:
                el = OxmlElement("w:" + edge); borders.append(el)
            for k, v in kw[edge].items():
                el.set(qn("w:" + k), str(v))

def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)

def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)

def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)

def keep_with_next(par, on=True):
    par.paragraph_format.keep_with_next = on

def add_field(run, instr):
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, t, f3):
        run._r.append(e)

FUNCS_RE = re.compile(r"^[A-Za-zÀ-ÿ]{2,}$")

def prep_tex(tex):
    return tex.replace(r"\qquad", r"\text{    }").replace(r"\quad", r"\text{  }").replace(r"\;", r"\text{ }").replace(r"\,", r"\text{ }")

def upright(om):
    """Groupes de 2 lettres ou plus (lim, inf, ln, min, leg, ILLIQ...) en romain ; découpe les runs mixtes."""
    SPLIT = re.compile(r"([A-Za-zÀ-ÿ]{2,})")
    for r in list(om.iter(f"{{{M_NS}}}r")):
        t = r.find(f"{{{M_NS}}}t")
        if t is None or not t.text or r.find(f"{{{M_NS}}}rPr") is not None:
            continue
        parts = [x for x in SPLIT.split(t.text) if x]
        if not any(SPLIT.fullmatch(x) for x in parts):
            continue
        parent = r.getparent(); idx = parent.index(r); parent.remove(r)
        for k, x in enumerate(parts):
            nr = etree.Element(f"{{{M_NS}}}r")
            if SPLIT.fullmatch(x):
                rpr = etree.SubElement(nr, f"{{{M_NS}}}rPr"); sty = etree.SubElement(rpr, f"{{{M_NS}}}sty")
                sty.set(f"{{{M_NS}}}val", "p")
            for child in r:
                if child.tag != f"{{{M_NS}}}t" and child.tag != f"{{{M_NS}}}rPr":
                    nr.append(copy.deepcopy(child))
            nt = etree.SubElement(nr, f"{{{M_NS}}}t"); nt.text = x
            if x != x.strip(): nt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            parent.insert(idx + k, nr)
    return om

def omml_inline(tex):
    om = upright(latex_to_omml(prep_tex(tex), display=False))
    if om.tag == f"{{{M_NS}}}oMathPara":
        om = om.find(f"{{{M_NS}}}oMath")
    return om

TOKEN = re.compile(r"(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*]+\*)")

def add_rich(par, text, font=None, size=None, color=None, italic=False, bold=False, math_size=None):
    """Ajoute du texte avec **gras**, *italique* et $maths$ (OMML en ligne)."""
    text = resolve_refs(text)
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("$") and tok.endswith("$") and len(tok) > 2:
            om = omml_inline(tok[1:-1])
            if math_size:
                for r in om.iter(f"{{{M_NS}}}r"):
                    rpr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(math_size * 2)))
                    rpr.append(sz); r.insert(0, rpr) if r.find(f"{{{M_NS}}}rPr") is None else r.insert(1, rpr)
            par._p.append(om)
            continue
        b, i = bold, italic
        if tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, i = tok[1:-1], True
        run = par.add_run(tok)
        run.bold, run.italic = b, i
        if font:
            run.font.name = font
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        if size: run.font.size = Pt(size)
        if color: run.font.color.rgb = RGBColor.from_string(color)
    return par

def ensure_numbering(doc):
    """Crée des définitions de puces et de listes numérotées dans numbering.xml."""
    numbering = doc.part.numbering_part.element
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    def abstract(aid, fmt, txt):
        a = etree.SubElement(numbering, qn("w:abstractNum")); a.set(qn("w:abstractNumId"), str(aid))
        lvl = etree.SubElement(a, qn("w:lvl")); lvl.set(qn("w:ilvl"), "0")
        e = etree.SubElement(lvl, qn("w:start")); e.set(qn("w:val"), "1")
        e = etree.SubElement(lvl, qn("w:numFmt")); e.set(qn("w:val"), fmt)
        e = etree.SubElement(lvl, qn("w:lvlText")); e.set(qn("w:val"), txt)
        e = etree.SubElement(lvl, qn("w:lvlJc")); e.set(qn("w:val"), "left")
        ppr = etree.SubElement(lvl, qn("w:pPr")); ind = etree.SubElement(ppr, qn("w:ind"))
        ind.set(qn("w:left"), "567"); ind.set(qn("w:hanging"), "283")
        if fmt == "bullet":
            rpr = etree.SubElement(lvl, qn("w:rPr")); rf = etree.SubElement(rpr, qn("w:rFonts"))
            rf.set(qn("w:ascii"), "Symbol"); rf.set(qn("w:hAnsi"), "Symbol"); rf.set(qn("w:hint"), "default")
        # les abstractNum doivent précéder les num
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(a)
        return a
    abstract(9001, "bullet", ""); abstract(9002, "decimal", "%1.")
    doc._num_counter = 9100
    def new_num(aid):
        doc._num_counter += 1
        n = etree.SubElement(numbering, qn("w:num")); n.set(qn("w:numId"), str(doc._num_counter))
        e = etree.SubElement(n, qn("w:abstractNumId")); e.set(qn("w:val"), str(aid))
        if aid == 9002:
            o = etree.SubElement(n, qn("w:lvlOverride")); o.set(qn("w:ilvl"), "0")
            s = etree.SubElement(o, qn("w:startOverride")); s.set(qn("w:val"), "1")
        return doc._num_counter
    doc._new_num = new_num

def set_num(par, num_id):
    pPr = par._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num_id))
    numPr.append(il); numPr.append(ni); pPr.append(numPr)


# ---------------------------------------------------------------------------
# moteur de rendu commun, paramétré par un « thème »
# ---------------------------------------------------------------------------
class Renderer:
    def __init__(self, doc, th):
        self.doc, self.th = doc, th
        self.n1 = self.n2 = self.n3 = 0
        self.ntab = 0; self.neq = 0; self.nfig = 0
        ensure_numbering(doc)

    def para(self, style=None, align=None, before=None, after=None, line=None, indent_first=None):
        p = self.doc.add_paragraph(style=style or self.th["body_style"])
        pf = p.paragraph_format
        if align is not None: p.alignment = align
        if before is not None: pf.space_before = Pt(before)
        if after is not None: pf.space_after = Pt(after)
        if line is not None: pf.line_spacing = line
        if indent_first is not None: pf.first_line_indent = Cm(indent_first)
        return p

    def text(self, t, **kw):
        th = self.th
        p = self.para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=th["after"], line=th["line"])
        add_rich(p, t, font=th["font"], size=th["size"], **kw)
        return p

    def heading(self, level, t, numbered=True):
        th = self.th
        if level == 1:
            self.n1 += 1; self.n2 = self.n3 = 0
            num = f"{self.n1}{th['h1_suffix']}" if numbered else ""
        elif level == 2:
            self.n2 += 1; self.n3 = 0
            num = f"{self.n1}.{self.n2}" if numbered else ""
        else:
            self.n3 += 1
            num = f"{self.n1}.{self.n2}.{self.n3}" if numbered else ""
        spec = th[f"h{level}"]
        p = self.para(style=spec.get("style"), align=WD_ALIGN_PARAGRAPH.LEFT, before=spec["before"], after=spec["after"])
        keep_with_next(p)
        pPr = p._p.get_or_add_pPr()
        ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), str(level - 1)); pPr.append(ol)
        label = (num + th["num_gap"] + t) if num else t
        add_rich(p, label, font=th["hfont"], size=spec["size"], color=spec.get("color"), bold=spec.get("bold", True), italic=spec.get("italic", False))
        if spec.get("smallcaps"):
            for r in p.runs: r.font.small_caps = True
        return p

    def bullets(self, items, ordered=False):
        th = self.th
        nid = self.doc._new_num(9002 if ordered else 9001)
        for it in items:
            p = self.para(style=th.get("list_style"), align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=th["after"] / 2, line=th["line"])
            set_num(p, nid)
            add_rich(p, it, font=th["font"], size=th["size"])

    def equation(self, tex, label):
        th = self.th
        self.neq += 1
        tbl = self.doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_borders(tbl)
        w1, w2 = th["text_width"] - 1.4, 1.4
        for row in tbl.rows:
            row.cells[0].width = Cm(w1); row.cells[1].width = Cm(w2)
        c0, c1 = tbl.rows[0].cells
        p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(3); p0.paragraph_format.space_after = Pt(3)
        om = upright(latex_to_omml(prep_tex(tex), display=True))
        if om.tag != f"{{{M_NS}}}oMathPara":
            para = etree.SubElement(p0._p, f"{{{M_NS}}}oMathPara"); para.append(om)
        else:
            p0._p.append(om)
        p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_rich(p1, f"({self.neq})", font=th["font"], size=th["size"])
        c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        # espace après
        sp = self.para(after=0, before=0); sp.paragraph_format.line_spacing = Pt(4)

    def figure(self, fname, caption, width):
        th = self.th
        self.nfig += 1
        cap = re.sub(r"^Figure \d+ – ", "", caption)
        path = os.path.join(FIG, fname + ".png")
        w = min(width, th["text_width"])
        if th["fig_caption_above"]:
            cp = self.para(style=th.get("fig_caption_style"), align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=4)
            keep_with_next(cp)
            if th.get("fig_caption_style"):
                add_rich(cp, " " + cap, font=th["font"], size=th["cap_size"])
            else:
                add_rich(cp, f"**Figure {self.nfig} –** " + cap, font=th["font"], size=th["cap_size"])
            p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
            p.add_run().add_picture(path, width=Cm(w))
        else:
            p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2)
            keep_with_next(p)
            p.add_run().add_picture(path, width=Cm(w))
            cp = self.para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
            add_rich(cp, f"**Figure {self.nfig}.** " + cap, font=th["font"], size=th["cap_size"])

    def table(self, head, rows, caption, widths=None, note=None, font=None):
        th = self.th
        self.ntab += 1
        cp = self.para(align=WD_ALIGN_PARAGRAPH.LEFT if th["fig_caption_above"] else WD_ALIGN_PARAGRAPH.JUSTIFY, before=8, after=4)
        keep_with_next(cp)
        add_rich(cp, f"**Tableau {self.ntab} {th['tab_sep']}** " + caption, font=th["font"], size=th["cap_size"])
        ncol = len(head)
        tbl = self.doc.add_table(rows=1 + len(rows), cols=ncol)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_borders(tbl); cell_margins(tbl)
        if widths is None:
            widths = [th["text_width"] / ncol] * ncol
        tot = sum(widths); widths = [w * th["text_width"] / tot for w in widths]
        fs = font or th["tab_size"]
        numre = re.compile(r"^[\s\-−–0-9,.%()/;\[\]pbM€kjLA-F→=*]*$")
        textcol = [j > 0 and not all(numre.match(str(r[j])) for r in rows) and sum(len(str(r[j])) for r in rows) / max(len(rows), 1) > 14 for j in range(ncol)]
        for i, row in enumerate([head] + rows):
            cells = tbl.rows[i].cells
            for j, val in enumerate(row):
                c = cells[j]; c.width = Cm(widths[j])
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 or textcol[j]) else WD_ALIGN_PARAGRAPH.CENTER
                if i < len(rows): p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                add_rich(p, str(val), font=th["font"], size=fs, bold=(i == 0), color=(th["tab_head_color"] if i == 0 else None))
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if i == 0:
                    if th["tab_head_fill"]: shade(c, th["tab_head_fill"])
                    set_cell_border(c, top={"val": "single", "sz": th["rule_thick"], "color": th["rule_color"]},
                                    bottom={"val": "single", "sz": th["rule_thin"], "color": th["rule_color"]})
                if i == len(rows):
                    set_cell_border(c, bottom={"val": "single", "sz": th["rule_thick"], "color": th["rule_color"]})
                elif i > 0 and th["row_rules"]:
                    set_cell_border(c, bottom={"val": "single", "sz": 2, "color": "D9D9D9"})
        # répéter l'en-tête et éviter les coupures de ligne
        trPr = tbl.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
        for r in tbl.rows:
            trPr = r._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        if note:
            npar = self.para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=3, after=8)
            add_rich(npar, "Note : " + note, font=th["font"], size=th["cap_size"] - 0.5, italic=True)
        else:
            self.para(after=4).paragraph_format.line_spacing = Pt(6)

    def render(self, blocks, numbered=True):
        for b in blocks:
            k = b[0]
            if k == "h1": self.heading(1, b[1], numbered)
            elif k == "h2": self.heading(2, b[1], numbered)
            elif k == "h3": self.heading(3, b[1], numbered)
            elif k == "p": self.text(b[1])
            elif k == "eq": self.equation(b[1], b[2])
            elif k == "bullets": self.bullets(b[1])
            elif k == "enum": self.bullets(b[1], ordered=True)
            elif k == "fig": self.figure(b[1], b[2], b[3])
            elif k == "table": self.table(b[1], b[2], b[3], b[4], b[5], b[6])
            elif k == "pagebreak": self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def references(self, title, refs, numbered_title=None):
        th = self.th
        self.heading(1, title, numbered=False)
        for r in sorted(refs, key=lambda s: s.lower()):
            p = self.para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3, line=1.0)
            p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6)
            add_rich(p, r, font=th["font"], size=th["ref_size"])


# ---------------------------------------------------------------------------
# 1. Format ArXiv
# ---------------------------------------------------------------------------
def build_arxiv(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = Cm(2.6); sec.bottom_margin = Cm(2.4)
    FONT = "Cambria"
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.12
    for name, size in (("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        hs = doc.styles[name]; hs.font.name = FONT; hs.font.size = Pt(size); hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
        rf = hs.element.rPr.find(qn("w:rFonts"))
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf is not None and rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        rf.set(qn("w:ascii"), FONT); rf.set(qn("w:hAnsi"), FONT)
    # langue française
    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "fr-FR"); rpr.append(lang)
    th = dict(font=FONT, hfont=FONT, size=10.5, after=6, line=1.12, body_style=None, list_style="List Paragraph",
              text_width=16.0, h1_suffix="", num_gap="   ",
              h1=dict(style="Heading 1", size=13, before=16, after=6), h2=dict(style="Heading 2", size=11.5, before=10, after=4),
              h3=dict(style="Heading 3", size=10.5, before=8, after=3, italic=True),
              fig_caption_above=False, cap_size=9.5, tab_size=9, tab_sep="–", tab_head_fill=None, tab_head_color=None,
              rule_thick=12, rule_thin=6, rule_color="000000", row_rules=False, ref_size=9.5)
    R = Renderer(doc, th)
    # en-tête « preprint » et pied de page
    hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_rich(hp, "Prépublication — " + C.META["title_short"], font=FONT, size=8, color="666666", italic=True)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(); run.font.size = Pt(9); run.font.name = FONT; add_field(run, "PAGE")
    # mention verticale de type arXiv (marge gauche) : simple paragraphe gris en tête de page 1
    p = R.para(align=WD_ALIGN_PARAGRAPH.LEFT, after=14)
    add_rich(p, "Prépublication au format arXiv — catégorie suggérée : q-fin.RM (Risk Management) — 15 octobre 2026", font=FONT, size=8.5, color="7F7F7F")
    # titre
    p = R.para(align=WD_ALIGN_PARAGRAPH.CENTER, after=14, line=1.05)
    add_rich(p, C.META["title"], font=FONT, size=17, bold=True)
    p = R.para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    for i, (a, _) in enumerate(C.META["authors"]):
        r = p.add_run(a); r.font.size = Pt(11.5); r.font.name = FONT; r.font.small_caps = True
        s = p.add_run(str(i + 1)); s.font.superscript = True; s.font.size = Pt(9)
        if i == 0: p.add_run("        ")
    for i, (_, aff) in enumerate(C.META["authors"]):
        q = R.para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0)
        add_rich(q, f"{i + 1} " + aff, font=FONT, size=8.5, color="404040", italic=True)
    q = R.para(align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=16)
    add_rich(q, C.META["version"], font=FONT, size=9, color="404040")
    # résumés
    for head, txt, kwh, kw in (("Résumé", C.META["abstract_fr"], "Mots-clés", C.META["keywords_fr"]),
                               ("Abstract", C.META["abstract_en"], "Keywords", C.META["keywords_en"])):
        h = R.para(align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=4)
        add_rich(h, head, font=FONT, size=10.5, bold=True)
        a = R.para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, line=1.05)
        a.paragraph_format.left_indent = Cm(1.0); a.paragraph_format.right_indent = Cm(1.0)
        add_rich(a, txt, font=FONT, size=9.5)
        k = R.para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10, line=1.0)
        k.paragraph_format.left_indent = Cm(1.0); k.paragraph_format.right_indent = Cm(1.0)
        add_rich(k, f"**{kwh} :** " + kw, font=FONT, size=9)
    k = R.para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4)
    k.paragraph_format.left_indent = Cm(1.0)
    add_rich(k, "**Classification JEL :** " + C.META["jel"] + ".   **Données :** historiques fictifs calibrés (annexe A) ; codes reproductibles.", font=FONT, size=9)
    R.render(C.B)
    R.references("Références", C.REFS)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    R.render(C.APP, numbered=False)
    doc.save(path)
    return R


# ---------------------------------------------------------------------------
# 2. Format BMI (gabarit Word de l'ancienne version : en-tête SAF, styles, notes de bas de page)
# ---------------------------------------------------------------------------
def build_bmi(path, template):
    doc = Document(template)
    body = doc.element.body
    paras = body.findall(qn("w:p"))
    keep = paras[:3]                      # titre, version, auteurs (avec appels de notes)
    for el in list(body):
        if el.tag == qn("w:sectPr") or el in keep:
            continue
        body.remove(el)
    def set_par_text(p_el, text):
        runs = p_el.findall(qn("w:r"))
        tr = [r for r in runs if r.find(qn("w:t")) is not None]
        tr[0].find(qn("w:t")).text = text
        for r in tr[1:]:
            p_el.remove(r)
    set_par_text(keep[0], C.META["title"])
    set_par_text(keep[1], C.META["version"])
    # notes de bas de page des auteurs
    fpart = None
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/footnotes"):
            fpart = rel.target_part
    if fpart is not None:
        fx = etree.fromstring(fpart.blob)
        for fn in fx.findall(qn("w:footnote")):
            fid = fn.get(qn("w:id"))
            if fid in ("2", "3"):
                ts = [t for t in fn.iter(qn("w:t"))]
                if ts:
                    full = C.META["authors"][0][1] if fid == "2" else C.META["authors"][1][1]
                    name = "Matthieu Moreau est actuaire et consultant." if fid == "2" else \
                           "Frédéric Planchet est professeur à l'ISFA et actuaire associé chez Prim'Act."
                    ts[0].text = name
                    for t in ts[1:]: t.text = ""
        fpart._blob = etree.tostring(fx, xml_declaration=True, encoding="UTF-8", standalone=True)
    FONT = "Candara"
    th = dict(font=FONT, hfont=FONT, size=11.5, after=6, line=1.0, body_style="Normal", list_style="List Paragraph",
              text_width=16.0, h1_suffix=".", num_gap="  ",
              h1=dict(size=14, before=24, after=12), h2=dict(size=12.5, before=14, after=6),
              h3=dict(size=11.5, before=10, after=4, italic=True),
              fig_caption_above=True, fig_caption_style="Figure", cap_size=10.5, tab_size=9.5, tab_sep=":",
              tab_head_fill="DCE6F1", tab_head_color="1F3864", rule_thick=8, rule_thin=4, rule_color="1F3864", row_rules=True,
              ref_size=10.5)
    R = Renderer(doc, th)
    # résumé (style du gabarit : Candara 10 pt, justifié)
    sp = R.para(after=6); sp = R.para(after=6)
    a = R.para(style="Corps", align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, line=1.0)
    add_rich(a, C.META["abstract_fr"], font=FONT, size=10)
    k = R.para(style="Corps", align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, line=1.0)
    add_rich(k, "**Mots-clés :** " + C.META["keywords_fr"] + ".", font=FONT, size=10)
    k = R.para(style="Corps", align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, line=1.0)
    add_rich(k, "**Classification JEL :** " + C.META["jel"] + ".", font=FONT, size=10)
    R.render(C.B)
    R.references("Références", C.REFS)
    R.render(C.APP, numbered=False)
    doc.save(path)
    return R


if __name__ == "__main__":
    ra = build_arxiv(os.path.join(OUT, "Article_Liquidite_ArXiv.docx"))
    rb = build_bmi(os.path.join(OUT, "Article_Liquidite_BMI.docx"), os.path.join(HERE, "..", "src", "bmi_old.docx"))
    print("équations :", ra.neq, "figures :", ra.nfig, "tableaux :", ra.ntab)
    print({k: v for k, v in EQN.items()})
